"""Client list used for the 'International client list' page of the About Company PDF.

The list lives in data/clients.xlsx (columns: Name of the client, Industry, Product,
Country of Origin, Continent). Replace that file on GitHub whenever the list changes.

Ranking for a new quotation (strongest first):
    1. same product   2. same industry   3. same country   4. same continent
"""

import os
import re

from openpyxl import load_workbook
from countries import country_key, continent_of

CLIENTS_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "data", "clients.xlsx")
MAX_CLIENTS = 38

_STOP = {"and", "of", "the", "for", "product", "products", "item", "items", "other", "others"}
_cache = {"mtime": None, "clients": []}


def _tokens(text):
    out = set()
    for word in re.split(r"[^a-z0-9]+", (text or "").lower()):
        if not word or word in _STOP:
            continue
        if word.endswith("ies") and len(word) > 4:
            word = word[:-3] + "y"
        elif word.endswith("s") and not word.endswith("ss") and len(word) > 3:
            word = word[:-1]
        out.add(word)
    return out


def load_clients():
    """Read data/clients.xlsx (re-read automatically if the file changes)."""
    mtime = os.path.getmtime(CLIENTS_FILE)
    if _cache["mtime"] == mtime:
        return _cache["clients"]
    ws = load_workbook(CLIENTS_FILE, data_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h or "").strip().lower() for h in rows[0]]

    def col(*names):
        for n in names:
            if n in header:
                return header.index(n)
        raise ValueError(f"clients.xlsx has no '{names[0]}' column")

    i_name, i_ind = col("name of the client", "name", "client"), col("industry")
    i_prod, i_ctry = col("product"), col("country of origin", "country")
    i_cont = col("continent")
    clients = []
    for r in rows[1:]:
        name = " ".join(str(r[i_name] or "").split())
        if not name:
            continue
        clients.append({
            "name": name,
            "industry": " ".join(str(r[i_ind] or "").split()),
            "product": " ".join(str(r[i_prod] or "").split()),
            "country": " ".join(str(r[i_ctry] or "").split()),
            "continent": " ".join(str(r[i_cont] or "").split()),
        })
    _cache.update(mtime=mtime, clients=clients)
    return clients


def industries():
    return sorted({c["industry"] for c in load_clients() if c["industry"]})


def rank_clients(product, industry, country, limit=MAX_CLIENTS):
    """Best-matching clients first. Returns a list of client dicts with a 'why' note."""
    wanted = _tokens(product.replace("&", " "))
    want_industry = (industry or "").strip().lower()
    want_country = country_key(country)
    want_continent = (continent_of(country) or "").lower()

    scored = []
    for idx, c in enumerate(load_clients()):
        by_product = bool(wanted & _tokens(c["product"]))
        by_industry = bool(want_industry) and c["industry"].lower() == want_industry
        by_country = bool(want_country) and country_key(c["country"]) == want_country
        by_continent = bool(want_continent) and c["continent"].lower() == want_continent
        key = (not by_product, not by_industry, not by_country, not by_continent, idx)
        why = [w for w, hit in (("product", by_product), ("industry", by_industry),
                                ("country", by_country), ("continent", by_continent)) if hit]
        scored.append((key, dict(c, why=", ".join(why) or "-")))
    scored.sort(key=lambda s: s[0])
    return [c for _, c in scored[:limit]]


def display_name(client):
    return f"{client['name']} ({client['country']})" if client["country"] else client["name"]
